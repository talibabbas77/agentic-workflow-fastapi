"""Turn uploaded or fetched files into titled sections of plain text.

Every format ends up as a list of sections with a heading path, so chunking and
citations work the same way for Markdown, HTML, PDF, and Word files.
"""

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field

import trafilatura
from docx import Document as DocxDocument
from docx.table import Table
from pypdf import PdfReader

MIME_BY_EXTENSION = {
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
TEXT_MIMES = {"text/markdown", "text/plain", "text/html", "application/xhtml+xml"}
DOCX = MIME_BY_EXTENSION[".docx"]

# Limits that keep one document from taking the worker down. A Word file is a zip,
# so a small upload can expand to gigabytes of XML; check before parsing anything.
MAX_ZIP_ENTRIES = 2_000
MAX_ZIP_UNCOMPRESSED = 50_000_000
MAX_ZIP_RATIO = 100
# What one document may produce: about 1,000 printed pages of text.
MAX_TEXT_CHARS = 2_000_000
MAX_BLOCKS = 100_000
MAX_PDF_PAGES = 2_000

_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_FENCE = re.compile(r"^\s*(```|~~~)")


class UnsupportedDocument(Exception):
    pass


@dataclass(frozen=True, slots=True)
class Section:
    headings: tuple[str, ...]
    text: str
    page: int | None = None


@dataclass(slots=True)
class Extracted:
    title: str
    sections: list[Section] = field(default_factory=list)
    page_count: int | None = None

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.sections)

    @property
    def content_hash(self) -> bytes:
        return hashlib.sha256(self.text.encode()).digest()


def detect_mime(filename: str, data: bytes) -> str:
    """Trust the bytes over the file name, and accept only formats we can read."""
    if data.startswith(b"%PDF-"):
        return "application/pdf"
    if data.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    check_archive(archive)
                    return DOCX
        except zipfile.BadZipFile:
            pass
        raise UnsupportedDocument("Only Word documents (.docx) are supported among zip files.")
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    mime = MIME_BY_EXTENSION.get(extension)
    if mime is None or mime not in TEXT_MIMES:
        raise UnsupportedDocument("Upload PDF, Word (.docx), Markdown, text, or HTML files.")
    if b"\x00" in data[:8192]:
        raise UnsupportedDocument("This file looks binary, not text.")
    return mime


def check_archive(archive: zipfile.ZipFile) -> None:
    """Refuse a zip that expands far beyond its size, before any of it is read."""
    entries = archive.infolist()
    if len(entries) > MAX_ZIP_ENTRIES:
        raise UnsupportedDocument("This Word document has too many parts to read.")
    total = 0
    for entry in entries:
        total += entry.file_size
        packed = max(entry.compress_size, 1)
        if entry.file_size > 1_000_000 and entry.file_size / packed > MAX_ZIP_RATIO:
            raise UnsupportedDocument("This Word document is compressed too much to read safely.")
    if total > MAX_ZIP_UNCOMPRESSED:
        raise UnsupportedDocument("This Word document is too large to read. Split it into parts.")


def extract(data: bytes, mime: str, fallback_title: str) -> Extracted:
    if mime == "application/pdf":
        result = _pdf(data, fallback_title)
    elif mime == DOCX:
        result = _docx(data, fallback_title)
    elif mime in ("text/html", "application/xhtml+xml"):
        result = _html(data, fallback_title)
    elif mime in ("text/markdown", "text/plain"):
        text = _decode(data)
        if len(text) > MAX_TEXT_CHARS:
            raise UnsupportedDocument("This file has too much text to read. Split it into parts.")
        result = markdown_sections(text, _first_heading(text) or fallback_title)
    else:
        raise UnsupportedDocument(f"Unsupported content type {mime}.")
    if not result.text.strip():
        raise UnsupportedDocument("No readable text found in this document.")
    return result


def markdown_sections(text: str, title: str) -> Extracted:
    """Split Markdown by headings, ignoring # lines inside code fences."""
    stack: list[tuple[int, str]] = []
    sections: list[Section] = []
    buffer: list[str] = []
    in_fence = False

    def flush() -> None:
        body = _clean("\n".join(buffer))
        if body:
            sections.append(Section(tuple(h for _, h in stack), body))
        buffer.clear()

    for line in text.splitlines():
        if _FENCE.match(line):
            in_fence = not in_fence
        match = None if in_fence else _HEADING.match(line)
        if match:
            flush()
            level = len(match.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, _clean_inline(match.group(2))))
        else:
            buffer.append(line)
    flush()
    return Extracted(title=title, sections=sections)


def _html(data: bytes, fallback_title: str) -> Extracted:
    html = _decode(data)
    markdown = trafilatura.extract(
        html,
        output_format="markdown",
        include_tables=True,
        include_links=False,
        include_images=False,
        include_comments=False,
        favor_recall=True,
    )
    metadata = trafilatura.extract_metadata(html)
    title = (metadata.title if metadata and metadata.title else "") or fallback_title
    return markdown_sections(markdown or "", title)


def _pdf(data: bytes, fallback_title: str) -> Extracted:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnsupportedDocument("This PDF is password protected.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise UnsupportedDocument(
                f"This PDF has more than {MAX_PDF_PAGES:,} pages. Split it into parts."
            )
        pages: list[str] = []
        size = 0
        for page in reader.pages:
            text = page.extract_text() or ""
            size += len(text)
            if size > MAX_TEXT_CHARS:
                raise UnsupportedDocument(
                    "This PDF has too much text to read. Split it into parts."
                )
            pages.append(text)
    except UnsupportedDocument:
        raise
    except Exception as exc:
        raise UnsupportedDocument("This PDF could not be read.") from exc
    info_title = (reader.metadata.title if reader.metadata else None) or ""
    title = info_title.strip() or fallback_title
    sections = [
        Section((title,), body, page=number)
        for number, raw in enumerate(pages, start=1)
        if (body := _clean(raw))
    ]
    return Extracted(title=title, sections=sections, page_count=len(pages))


def _docx(data: bytes, fallback_title: str) -> Extracted:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            check_archive(archive)
        document = DocxDocument(io.BytesIO(data))
    except UnsupportedDocument:
        raise
    except Exception as exc:
        raise UnsupportedDocument("This Word document could not be read.") from exc
    # Resolve each style id once; looking it up per paragraph is slow on long files.
    names = {style.style_id: style.name or "" for style in document.styles}
    lines: list[str] = []
    size = 0
    title = (document.core_properties.title or "").strip()
    for count, block in enumerate(document.iter_inner_content()):  # reading order
        if count >= MAX_BLOCKS or size > MAX_TEXT_CHARS:
            raise UnsupportedDocument(
                "This Word document is too long to read. Split it into parts."
            )
        if isinstance(block, Table):
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in block.rows]
            size += sum(len(row) for row in rows)
            lines.extend(["", *rows, ""])
            continue
        text = block.text.strip()
        size += len(text)
        style = names.get(block._p.style or "", "")
        if style == "Title" and text:
            title = title or text
            continue
        level = _heading_level(style)
        lines.append(f"{'#' * level} {text}" if level and text else text)
    return markdown_sections("\n".join(lines), title or fallback_title)


def _heading_level(style: str) -> int:
    match = re.match(r"Heading (\d)", style)
    return min(int(match.group(1)), 6) if match else 0


def _first_heading(text: str) -> str | None:
    for line in text.splitlines()[:50]:
        match = _HEADING.match(line)
        if match and len(match.group(1)) == 1:
            return _clean_inline(match.group(2))
    return None


def _decode(data: bytes) -> str:
    if len(data) > MAX_TEXT_CHARS * 4:
        raise UnsupportedDocument("This file has too much text to read. Split it into parts.")
    return _decode_text(data)


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _clean_inline(text: str) -> str:
    return re.sub(r"[*_`]+", "", text).strip()
