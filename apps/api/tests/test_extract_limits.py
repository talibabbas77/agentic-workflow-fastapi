"""Hostile documents fail one upload, never the worker (security review, 2026-10-14)."""

import io
import zipfile

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from docent.knowledge import extract as extract_module
from docent.knowledge.extract import DOCX, UnsupportedDocument, detect_mime, extract
from docent.knowledge.isolated import extract_isolated


def docx_bytes(paragraphs: list[str]) -> bytes:
    document = DocxDocument()
    document.add_heading("Refunds", level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def zip_bomb() -> bytes:
    """A real .docx whose document.xml is two million empty paragraphs."""
    source = zipfile.ZipFile(io.BytesIO(docx_bytes(["x"])))
    body = b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
    body += b"<w:p/>" * 2_000_000 + b"</w:body></w:document>"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as bomb:
        for item in source.infolist():
            data = body if item.filename == "word/document.xml" else source.read(item.filename)
            bomb.writestr(item.filename, data)
    return out.getvalue()


def test_a_word_zip_bomb_is_refused_at_upload_and_at_parse() -> None:
    bomb = zip_bomb()
    assert len(bomb) < 100_000
    with pytest.raises(UnsupportedDocument, match="compressed too much"):
        detect_mime("handbook.docx", bomb)
    with pytest.raises(UnsupportedDocument, match="compressed too much"):
        extract(bomb, DOCX, "handbook")


def test_long_documents_stop_at_the_caps(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extract_module, "MAX_BLOCKS", 50)
    with pytest.raises(UnsupportedDocument, match="too long"):
        extract(docx_bytes([f"Line {n}" for n in range(100)]), DOCX, "handbook")

    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=200, height=200)
    pdf = io.BytesIO()
    writer.write(pdf)
    monkeypatch.setattr(extract_module, "MAX_PDF_PAGES", 2)
    with pytest.raises(UnsupportedDocument, match="more than 2 pages"):
        extract(pdf.getvalue(), "application/pdf", "guide")

    monkeypatch.setattr(extract_module, "MAX_TEXT_CHARS", 1_000)
    with pytest.raises(UnsupportedDocument, match="too much text"):
        extract(b"# Notes\n\n" + b"a" * 2_000, "text/markdown", "notes")


async def test_word_and_pdf_files_are_read_in_a_separate_process() -> None:
    data = docx_bytes(["Refunds take five business days."])
    extracted = await extract_isolated(data, DOCX, "fallback", seconds=60, memory_mb=1024)
    assert extracted.sections[0].headings == ("Refunds",)
    assert extracted.sections[0].text == "Refunds take five business days."

    with pytest.raises(UnsupportedDocument, match="took too long"):
        await extract_isolated(data, DOCX, "fallback", seconds=0.01, memory_mb=1024)

    bomb = zip_bomb()
    with pytest.raises(UnsupportedDocument, match="compressed too much"):
        await extract_isolated(bomb, DOCX, "fallback", seconds=60, memory_mb=1024)
