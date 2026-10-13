"""Parsing and chunking."""

import io

import pytest
from docx import Document

from docent.knowledge.chunking import MAX_CHARS, TARGET_CHARS, chunk
from docent.knowledge.extract import (
    DOCX,
    Section,
    UnsupportedDocument,
    detect_mime,
    extract,
    markdown_sections,
)


def make_pdf(text: str) -> bytes:
    """A valid one-page PDF that draws one line of text, with a correct xref table."""
    stream = f"BT /F1 18 Tf 20 60 Td ({text}) Tj ET".encode()
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n%s\nendstream" % (len(stream), stream),
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    return bytes(out)


MINIMAL_PDF = make_pdf("Refunds take 30 days")


def test_markdown_headings_build_paths_and_ignore_code_fences() -> None:
    text = """# Guide

Intro text.

## Setup

```bash
# not a heading
pip install thing
```

### Linux

Use the package.

## Billing

Pay monthly.
"""
    result = markdown_sections(text, "Guide")
    paths = [(s.headings, s.text.splitlines()[0]) for s in result.sections]
    assert paths == [
        (("Guide",), "Intro text."),
        (("Guide", "Setup"), "```bash"),
        (("Guide", "Setup", "Linux"), "Use the package."),
        (("Guide", "Billing"), "Pay monthly."),
    ]


def test_detect_mime_trusts_content_over_extension() -> None:
    assert detect_mime("notes.md", b"# Hi") == "text/markdown"
    assert detect_mime("report.txt", MINIMAL_PDF) == "application/pdf"
    with pytest.raises(UnsupportedDocument):
        detect_mime("image.png", b"\x89PNG\r\n")
    with pytest.raises(UnsupportedDocument):
        detect_mime("archive.zip", b"PK\x03\x04garbage")


def test_pdf_text_and_pages() -> None:
    result = extract(MINIMAL_PDF, "application/pdf", "policy")
    assert result.page_count == 1
    assert result.sections[0].page == 1
    assert "Refunds take 30 days" in result.text


def test_docx_headings_paragraphs_and_tables_in_order() -> None:
    document = Document()
    document.add_heading("Returns", level=1)
    document.add_paragraph("Ship items back within 14 days.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Days"
    table.cell(1, 0).text, table.cell(1, 1).text = "EU", "14"
    document.add_heading("Refunds", level=1)
    document.add_paragraph("Money returns in 5 days.")
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()

    assert detect_mime("returns.docx", data) == DOCX
    result = extract(data, DOCX, "returns")
    assert [s.headings for s in result.sections] == [("Returns",), ("Refunds",)]
    assert "EU | 14" in result.sections[0].text


def test_html_keeps_main_content_only() -> None:
    html = b"""<html><head><title>Plans</title></head><body>
    <nav>Home Pricing Blog Careers</nav>
    <main><h1>Plans</h1>
    <p>The starter plan includes three seats and email support for teams.</p>
    <h2>Upgrades</h2>
    <p>Upgrade at any time from the billing page in your account settings.</p></main>
    <footer>Copyright Acme</footer></body></html>"""
    result = extract(html, "text/html", "fallback")
    assert result.title == "Plans"
    assert "three seats" in result.text
    assert "Careers" not in result.text


def test_chunks_respect_sections_and_size_with_overlap() -> None:
    long_text = "\n\n".join(
        f"Sentence number {i} explains a detail in full." * 3 for i in range(60)
    )
    sections = [Section(("Guide", "Long"), long_text), Section(("Guide", "Short"), "Tiny note.")]
    passages = chunk(sections)
    assert passages[-1].heading_path == "Guide > Short"
    assert passages[-1].content == "Tiny note."
    long_parts = [p for p in passages if p.heading_path == "Guide > Long"]
    assert len(long_parts) > 3
    assert all(len(p.content) <= MAX_CHARS for p in passages)
    assert all(len(p.content) >= TARGET_CHARS // 2 for p in long_parts[:-1])
    # consecutive chunks share some text, so a sentence near a break keeps context
    assert long_parts[1].content.split("\n\n")[0] in long_parts[0].content
    assert [p.ordinal for p in passages] == list(range(len(passages)))


def test_oversized_paragraph_is_split_without_losing_text() -> None:
    word = "lorem "
    text = word * 1500  # one paragraph, no sentence breaks
    passages = chunk([Section(("Doc",), text.strip())])
    assert all(len(p.content) <= MAX_CHARS for p in passages)
    rebuilt = " ".join(p.content for p in passages).split()
    assert rebuilt.count("lorem") >= 1500
